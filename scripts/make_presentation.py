"""Презентация для защиты: docs/RepairCalc.pptx. Цифры берутся из reports/metrics.json и живого сервиса.

Запуск: python -m scripts.make_presentation
"""
import json
import sys
import time
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

GREEN = RGBColor(0x1F, 0x8A, 0x4C)
DARK = RGBColor(0x1D, 0x23, 0x30)
MUTED = RGBColor(0x6B, 0x73, 0x85)
LIGHT = RGBColor(0xE8, 0xF5, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W, H = Emu(12192000), Emu(6858000)  # 16:9
IMG = ROOT / "docs" / "img"


def cm(x):
    return Emu(int(x * 360000))


class Deck:
    def __init__(self):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = W, H
        self.n = 0

    def slide(self, title, notes="", kicker=None):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.n += 1
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, cm(0.35), H)
        bar.fill.solid(); bar.fill.fore_color.rgb = GREEN; bar.line.fill.background()
        if kicker:
            self.text(s, kicker, 1.5, 0.8, 30, 0.8, 12, MUTED)
        self.text(s, title, 1.5, 1.3, 30, 1.6, 28, DARK, bold=True)
        self.text(s, f"RepairCalc · {self.n}", 28.5, 18, 5, 0.6, 10, MUTED, align=PP_ALIGN.RIGHT)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def text(self, s, txt, x, y, w, h, size=16, color=DARK, bold=False, align=None):
        tb = s.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = txt if isinstance(txt, list) else [txt]
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = para.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.bold = bold
            if align:
                para.alignment = align
            para.space_after = Pt(6)
        return tb

    def bullets(self, s, items, x=1.5, y=3.4, w=30, size=17):
        return self.text(s, [f"•  {i}" for i in items], x, y, w, 14, size)

    def card(self, s, x, y, w, h, big, small, fill=LIGHT):
        r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cm(x), cm(y), cm(w), cm(h))
        r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background()
        r.adjustments[0] = 0.08
        self.text(s, big, x + 0.4, y + 0.3, w - 0.8, 1.4, 26, GREEN, bold=True)
        self.text(s, small, x + 0.4, y + 1.7, w - 0.8, h - 1.8, 13, DARK)

    def table(self, s, rows, x, y, w, col_w=None, size=12, header=True):
        t = s.shapes.add_table(len(rows), len(rows[0]), cm(x), cm(y), cm(w), cm(0.8 * len(rows))).table
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.text = str(val)
                para = c.text_frame.paragraphs[0]
                para.runs[0].font.size = Pt(size)
                para.runs[0].font.bold = header and i == 0
                para.runs[0].font.color.rgb = WHITE if (header and i == 0) else DARK
                c.fill.solid()
                c.fill.fore_color.rgb = GREEN if (header and i == 0) else (LIGHT if i % 2 == 0 else WHITE)
        if col_w:
            for j, cw in enumerate(col_w):
                t.columns[j].width = cm(cw)
        return t

    def image(self, s, name, x, y, w=None, h=None):
        f = IMG / name
        if f.exists():
            s.shapes.add_picture(str(f), cm(x), cm(y), cm(w) if w else None, cm(h) if h else None)


def main():
    m = json.loads((ROOT / "reports" / "metrics.json").read_text())
    hyb, base = m["results"]["hybrid"], m["results"]["baseline"]
    ds = m["dataset"]
    from app.calc.estimate import estimate
    from app.service import get_index
    t0 = time.perf_counter()
    idx = get_index()
    start_s = time.perf_counter() - t0
    prods = list(idx.catalog)
    auto_share = sum(1 for p in prods if p["_wt"] and p["_source"] == "ml") / len(prods)
    demo = estimate(idx, "apartment", "capital", [{"type": "bathroom", "area": 6, "humidity": "high"}])
    f = lambda v: f"{v:.2f}".replace(".", ",")  # noqa: E731

    d = Deck()
    # 1
    s = d.slide("RepairCalc", kicker="SpbTechRun 2026 · кейс №1")
    d.text(s, "Микросервис умных калькуляторов ремонта: ML-подбор материалов под помещение "
              "и точный расчёт количества и стоимости по реальному каталогу", 1.5, 3.4, 26, 3, 20, DARK)
    d.text(s, "Команда «___» · капитан ___", 1.5, 15.5, 20, 1, 14, MUTED)
    d.image(s, "1_calc.png", 18, 7, w=15)

    # 2
    s = d.slide("Проблема: покупатель не знает, сколько и что купить", kicker="Боль",
                notes="Пользователь маркетплейса не понимает объём материалов и бюджет — бросает корзину, звонит в поддержку.")
    d.bullets(s, ["Сколько упаковок плитки на стену 3,2 × 2,7? Какой клей и сколько затирки?",
                  "Что можно в санузел, а что нельзя: обычный ГКЛ или ГКЛВ, какая краска?",
                  "Итог — брошенные корзины, звонки менеджерам, низкая конверсия «под задачу»",
                  "Готовые калькуляторы статичны, не учитывают влажность и назначение комнаты,",
                  "не привязаны к каталогу и остаткам и не масштабируются на тысячи SKU без ручной настройки"])

    # 3
    s = d.slide("Решение: три сценария на одном движке", kicker="Продукт",
                notes="Все три сценария используют общий движок: ML-разметка каталога + планировщик работ + формулы расхода.")
    d.card(s, 1.5, 3.6, 10, 6, "Смета", ["Квартира / студия / дом,", "косметический / капитальный /", "дизайнерский ремонт,",
                                          "помещения → товары, количество, итог"])
    d.card(s, 12.2, 3.6, 10, 6, "Комната", ["Влажность, назначение, нагрузка", "→ влагостойкие, моющиеся,", "износостойкие материалы",
                                            "+ пояснение «почему» (RAG)"])
    d.card(s, 22.9, 3.6, 10, 6, "Карточка", ["/api/v1/product/{id}/calc", "размеры стены → упаковки и цена", "+ клей, затирка, грунт", "одной кнопкой в корзину"])
    d.text(s, "В основе — ML, который сам относит тысячи товаров каталога к видам работ, классам материалов и помещениям",
           1.5, 11, 31, 2, 17, GREEN, bold=True)

    # 4
    s = d.slide(f"Демо: санузел 6 м², капремонт → {float(demo['total']):,.0f} ₽".replace(",", " "), kicker="Сценарий 1",
                notes="Показываем: ГКЛВ, гидроизоляция с заведением на стены, плитка с подрезкой, IP44, вентилятор, тёплый пол, итог и пояснение.")
    d.image(s, "1_calc.png", 1.5, 3.3, h=14)
    d.bullets(s, [f"{demo['lines_count']} позиций за {demo['elapsed_ms']:.0f} мс",
                  "Только влагостойкие материалы: ГКЛВ, плитка, IP44",
                  "Гидроизоляция пола + 30 см на стены, 2 слоя",
                  "Формула у каждой позиции",
                  "Альтернативы и пояснение «почему»"], x=22, y=3.6, w=11.5, size=14)

    # 5
    s = d.slide("Расчёт прямо в карточке товара", kicker="Сценарий 3",
                notes="Виджет встраивается в любую страницу (в т.ч. Bitrix) двумя строками. Сопутствующие товары — тот же вид работ.")
    d.image(s, "2_product.png", 1.5, 3.3, h=13.5)
    d.image(s, "4_widget.png", 19.5, 3.3, w=14)
    d.text(s, ["<div data-repaircalc-product=\"ID\"></div>", "<script src=\"…/ui/widget.js\"></script>"], 19.5, 13.2, 14, 2, 12, MUTED)

    # 6
    s = d.slide("ML-матчинг «товар → виды работ»", kicker="Где ML",
                notes="Гибрид: логрегрессия точна на знакомых типах; zero-shot по справочнику и kNN по эмбеддингам переносят знание на новые; шлюз отсекает товары вне ремонта.")
    d.table(s, [["Компонент", "Что даёт"],
                ["TF-IDF слов и символов + логрегрессия OvR", "точность на знакомых типах товаров"],
                ["Эмбеддинги multilingual-e5 (локально)", "смысл, а не буквальные совпадения"],
                ["Zero-shot по справочнику: сходство с типовыми материалами", "новые виды товаров без обучения"],
                ["kNN по размеченным товарам", "перенос разметки похожих товаров"],
                ["Шлюз «относится к ремонту»", "сад, быт, фасад → без меток"],
                ["Пороги по каждому виду работ (out-of-fold)", "баланс точности и полноты"],
                ["Второй классификатор: класс материала (48 классов)", "выбор формулы расхода"]],
            1.5, 3.4, 31, col_w=[17, 14], size=13)
    d.text(s, "Многометочная задача: 23% товаров относятся к нескольким видам работ", 1.5, 11.2, 31, 1, 15, GREEN, bold=True)

    # 7
    s = d.slide("Справочник: 26 видов работ, 5 этапов, 7 правил разметки", kicker="Таксономия",
                notes="Таксономия по технологической последовательности. Правила R1–R7 и разбор неоднозначных товаров — в docs/taxonomy.md. Отправлено организаторам на согласование.")
    d.bullets(s, ["Подготовка → черновые → инженерные системы → чистовая отделка → завершающие",
                  "Вид работ = законченная технологическая операция: «укладка плитки», «монтаж ГКЛ»",
                  "Товар → все работы, где он основной материал, расходник или специализированный инструмент",
                  "Неоднозначные: грунтовка (6 работ), монтажная пена, обои под покраску, полотенцесушитель",
                  "Вне охвата: кладка, фасад, сад, быт → пустой набор меток",
                  "Каждый вид работ связан с формулами калькулятора — разметка сразу работает в смете",
                  f"Эталон: {ds['products']} товаров, {ds['subtypes']} подтипов; 248 — на двойную ручную проверку (каппа Коэна)"])

    # 8
    gs, nw, rnd = hyb["group_5fold_unseen_subtypes"], hyb["new_products"], hyb["random_5fold"]
    s = d.slide("Качество: micro-F1 ≥ 0,8 и честная проверка обобщения", kicker="Метрики",
                notes="Случайное разбиение и новые товары — цель достигнута. Самый жёсткий протокол — незнакомые подтипы: модель не видела ни одного товара такого вида. Там zero-shot даёт основной прирост; дальше — LLM для неуверенных товаров.")
    d.table(s, [["Протокол", "micro-F1", "macro-F1", "nDCG@10", "вне ремонта"],
                ["5-fold по товарам", f(rnd["micro_f1"]), f(rnd["macro_f1"]), f(rnd["ndcg@10"]), f(rnd.get("out_of_scope_correct", 0))],
                ["Новые товары (отдельная выгрузка)", f(nw["micro_f1"]), f(nw["macro_f1"]), f(nw["ndcg@10"]), f(nw.get("out_of_scope_correct", 0))],
                ["Незнакомые подтипы товаров", f(gs["micro_f1"]), f(gs["macro_f1"]), f(gs["ndcg@10"]), f(gs.get("out_of_scope_correct", 0))],
                ["baseline TF-IDF, незнакомые подтипы", f(base["group_5fold_unseen_subtypes"]["micro_f1"]),
                 f(base["group_5fold_unseen_subtypes"]["macro_f1"]), f(base["group_5fold_unseen_subtypes"]["ndcg@10"]),
                 f(base["group_5fold_unseen_subtypes"].get("out_of_scope_correct", 0))]],
            1.5, 3.4, 31, col_w=[13, 4.5, 4.5, 4.5, 4.5], size=14)
    d.text(s, ["Каталог синтетический, поэтому на знакомых типах цифры завышены — основной ориентир для нас",
               "протокол «незнакомые подтипы»: гибрид против baseline даёт "
               f"+{(gs['micro_f1'] - base['group_5fold_unseen_subtypes']['micro_f1']) * 100:.0f} п.п. micro-F1 "
               f"и +{(gs['ndcg@10'] - base['group_5fold_unseen_subtypes']['ndcg@10']) * 100:.0f} п.п. nDCG@10. Всё воспроизводится: make eval"],
           1.5, 9.4, 31, 2.5, 14, MUTED)

    # 9
    s = d.slide("Математическая модель расчёта", kicker="Формулы",
                notes="Каждая позиция сметы показывает формулу. На 16 эталонных расчётах, сделанных вручную, MAE = 0 (tests/test_formulas.py).")
    d.table(s, [["Материал", "Формула"],
                ["Плитка, ламинат", "⌈S · (1 + запас) / S упаковки⌉; запас 10% / 15% диагональ / +5% при S < 4 м²"],
                ["Обои", "полосы = ⌈L стен / ширина⌉; в рулоне = ⌊длина / (h + раппорт + 0,1)⌋; рулоны = ⌈полосы / в рулоне⌉"],
                ["Краска, грунт", "⌈S · расход л/м² · слои / объём банки⌉"],
                ["Смеси (клей, штукатурка, стяжка, наливной пол)", "⌈S · расход кг/м²·мм · толщина / вес мешка⌉"],
                ["Гидроизоляция", "(S пола + P · 0,3 м) · 1,2 кг/м² · 2 слоя"],
                ["Линолеум", "полосы по ширине рулона × (длина + 0,1 м)"],
                ["Электрика", "кабель = точки · 7 м + 10 м; розетки по площади и назначению"]],
            1.5, 3.4, 31, col_w=[9, 22], size=12)
    d.text(s, "Расход берётся из характеристик карточки; если её нет — норма из базы знаний, с пометкой в смете. "
              "Учёт единицы продажи (упак., м², м, шт); деньги в Decimal; MAE на 16 эталонных расчётах = 0",
           1.5, 11.3, 31, 2, 14, GREEN, bold=True)

    # 10
    s = d.slide("Подбор под помещение и ранжирование", kicker="Рекомендации",
                notes="Жёсткие ограничения помещения + мягкие предпочтения. При отсутствии подходящих товаров — предупреждение, а не тихая замена.")
    d.bullets(s, ["Кандидаты: ML (или админка) отнёс товар к нужному виду работ и классу материала",
                  "Жёсткие ограничения: санузел → влагостойкие; балкон → морозостойкие; кухня, прихожая, детская → ламинат ≥ 33 класса",
                  "Ранжирование: 0,25·уверенность ML + 0,25·соответствие позиции + 0,15·популярность (recsys_rating)",
                  "+ 0,05·рейтинг + 0,15·наличие в регионе + 0,15·ценовой сегмент под тип ремонта",
                  "Детская → бонус экологичным, кухня → моющимся",
                  "Оптимизация фасовки: среди близких по рангу — минимальная итоговая стоимость",
                  "Пояснение RAG: база норм и требований к помещениям + локальная LLM (без неё — шаблон)"], size=16)

    # 11
    s = d.slide("Полный цикл: загрузка каталога → авторазметка → админка", kicker="Админка",
                notes="Синхронизация обходит дерево категорий, читает страницы до last_page, дочитывает полные карточки и размечает товары.")
    d.image(s, "3_admin.png", 1.5, 3.3, h=13.5)
    d.bullets(s, [f"{auto_share * 100:.1f}% каталога размечено автоматически",
                  "155 новых товаров со стенда: загрузка и разметка ≈ 1 с",
                  "Новый товар → вид работ, класс, помещения, похожие товары",
                  "Ручная привязка важнее ML, откат в один клик",
                  "Фильтр «низкая уверенность» → очередь на проверку"], x=21.5, y=3.6, w=12, size=14)

    # 12
    r = m["resources"]
    s = d.slide("Производительность и закрытый контур", kicker="Инфраструктура",
                notes="Все модели локальные. На сервере заказчика RTX 3090 отдаётся под LLM, эмбеддинги работают на CPU.")
    d.card(s, 1.5, 3.6, 7.5, 5, f"{hyb['timing']['predict_ms_per_product']:.1f} мс", ["ML-разметка товара"])
    d.card(s, 9.6, 3.6, 7.5, 5, f"{demo['elapsed_ms']:.0f} мс", ["смета на санузел", f"({demo['lines_count']} позиций)"])
    d.card(s, 17.7, 3.6, 7.5, 5, f"{hyb['timing']['train_s']:.0f} с", ["обучение модели", f"на {ds['products']} товарах"])
    d.card(s, 25.8, 3.6, 7.5, 5, f"{start_s:.0f} с", ["старт сервиса", "с разметкой каталога"])
    d.bullets(s, ["Эмбеддинги multilingual-e5-small (MIT), 470 МБ, работают на CPU",
                  "LLM: Ollama + Qwen2.5-7B/14B-Instruct (Apache-2.0) на RTX 3090; без LLM всё работает на правилах",
                  f"Пиковая память при обучении и оценке ≈ {r['peak_rss_mb'] / 1024:.1f} ГБ; Docker + docker-compose с GPU для Ollama",
                  "Лицензии проверены скриптом: GPL/AGPL нет (make licenses)"], y=9.6, size=15)

    # 13
    s = d.slide("Ценность и развитие", kicker="Бизнес",
                notes="Метрики эффекта для пилота: конверсия карточки в корзину, средний чек, доля корзин «под задачу», нагрузка на поддержку.")
    d.bullets(s, ["Рост конверсии и среднего чека: продажа «под задачу», сопутствующие товары одной кнопкой",
                  "Меньше ручной работы: авторазметка каталога, ручная — только для неуверенных товаров",
                  "Меньше звонков в поддержку: смета и пояснения прямо на сайте",
                  "Пилот: A/B на карточках плитки, ламината и обоев; метрики — CR в корзину, средний чек, время до корзины",
                  "Развитие: реальные данные заказов → обучаемые коэффициенты отходов; LLM для товаров с низкой уверенностью;",
                  "CV по плану квартиры; новые вертикали (сад, инструмент) = новый справочник + make train"])

    # 14
    s = d.slide("Спасибо! Живое демо", kicker="RepairCalc")
    d.bullets(s, ["UI: /ui/  ·  Swagger: /docs  ·  метрики: вкладка «Качество ML»",
                  "Код, справочник, методика, эталон, отчёт метрик — в репозитории",
                  "make setup model data train eval test run — воспроизводится с нуля"], size=18)

    out = ROOT / "docs" / "RepairCalc.pptx"
    d.p.save(out)
    print(f"{d.n} слайдов -> {out}")


if __name__ == "__main__":
    main()
